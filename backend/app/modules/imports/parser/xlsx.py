from pathlib import Path

from openpyxl import load_workbook

from app.modules.imports.parser.base import (
    SpreadsheetError,
    SpreadsheetMetadata,
    SpreadsheetPreview,
    ensure_not_empty,
    headers_from_values,
    normalize_cell,
)


XLSX_SIGNATURE = b"PK"


class XlsxSpreadsheetReader:
    file_type = "xlsx"

    def __init__(self, path: Path) -> None:
        self.path = path
        if self.path.read_bytes()[:2] != XLSX_SIGNATURE:
            raise SpreadsheetError(code="IMPORT_INVALID_FORMAT", message="File is not a valid XLSX archive")

    def _workbook(self):
        try:
            return load_workbook(self.path, read_only=True, data_only=True)
        except Exception as exc:
            raise SpreadsheetError(code="IMPORT_INVALID_FORMAT", message="Cannot read XLSX file") from exc

    def read_metadata(self) -> SpreadsheetMetadata:
        workbook = self._workbook()
        try:
            return SpreadsheetMetadata(sheet_names=list(workbook.sheetnames), file_type=self.file_type)
        finally:
            workbook.close()

    def preview(self, *, sheet_name: str | None = None, header_row: int = 1, limit: int = 20) -> SpreadsheetPreview:
        workbook = self._workbook()
        try:
            sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
            headers = headers_from_values([cell for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), [])])
            rows: list[list[object | None]] = []
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                normalized = [normalize_cell(cell) for cell in row[: len(headers)]]
                if any(cell is not None for cell in normalized):
                    rows.append(normalized)
                if len(rows) >= limit:
                    break
            ensure_not_empty(headers, rows)
            return SpreadsheetPreview(
                sheet_names=list(workbook.sheetnames),
                sheet=sheet.title,
                headers=headers,
                rows=rows,
                total_rows=max(sheet.max_row - header_row, 0),
                file_type=self.file_type,
            )
        finally:
            workbook.close()

    def iter_rows(self, *, sheet_name: str | None = None, header_row: int = 1) -> list[tuple[int, dict[str, object | None]]]:
        workbook = self._workbook()
        try:
            sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
            headers = headers_from_values([cell for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), [])])
            ensure_not_empty(headers)
            result: list[tuple[int, dict[str, object | None]]] = []
            for index, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
                normalized = [normalize_cell(cell) for cell in row[: len(headers)]]
                if any(cell is not None for cell in normalized):
                    result.append((index, dict(zip(headers, normalized, strict=False))))
            return result
        finally:
            workbook.close()
