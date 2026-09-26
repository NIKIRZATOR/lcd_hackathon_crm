"""Fixture adapters.  They only parse and normalize; matching lives in the service layer."""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SOURCES = {"WEBSITE", "LMS", "PAYMENT", "B2C_USER", "VENDOR_CATALOG"}
TEAM_FIXTURE_PATH = Path(__file__).parent / "fixtures" / "program_signals.json"


@dataclass(frozen=True)
class AdapterRecord:
    source: str
    external_key: str
    raw_payload: dict[str, Any]
    normalized_payload: dict[str, Any]


def normalize_text(value: object | None) -> str | None:
    if value is None:
        return None
    result = " ".join(str(value).replace("«", "").replace("»", "").replace('"', "").split())
    return result or None


def normalized_key(value: object | None) -> str | None:
    text = normalize_text(value)
    return text.casefold() if text else None


def normalize_email(value: object | None) -> str | None:
    text = normalize_text(value)
    return text.casefold() if text else None


def normalize_phone(value: object | None) -> str | None:
    digits = re.sub(r"\D", "", str(value or ""))
    if len(digits) == 11 and digits.startswith("8"):
        digits = "7" + digits[1:]
    return digits or None


def hash_value(value: str | None) -> str | None:
    return hashlib.sha256(value.encode("utf-8")).hexdigest() if value else None


def split_products(value: object | None) -> list[str]:
    text = normalize_text(value)
    if not text:
        return []
    return [part for part in (normalize_text(part) for part in re.split(r"\s*,\s*", text)) if part]


def _sheet_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    headers = [normalize_text(value) for value in next(iterator, ())]
    rows: list[dict[str, Any]] = []
    for values in iterator:
        row = {header: value for header, value in zip(headers, values) if header}
        if any(value is not None and str(value).strip() for value in row.values()):
            rows.append(row)
    return rows


def parse_vendor_fixture(path: Path) -> list[AdapterRecord]:
    rows = _sheet_rows(path)
    records: list[AdapterRecord] = []
    for row in rows:
        vendor = normalize_text(row.get("Компания"))
        for product in split_products(row.get("Продукт")):
            if not vendor:
                continue
            contact_name = normalize_text(row.get("ФИО"))
            external_key = ":".join(filter(None, [normalized_key(vendor), normalized_key(product), normalized_key(contact_name)]))
            records.append(AdapterRecord(
                source="VENDOR_CATALOG", external_key=external_key, raw_payload=row,
                normalized_payload={
                    "vendor_name": vendor, "vendor_key": normalized_key(vendor),
                    "product_name": product, "product_key": normalized_key(product),
                    "contact_name": contact_name, "phone": normalize_phone(row.get("Телефон")),
                    "email": normalize_email(row.get("Почта")), "preferred_channel": normalize_text(row.get("Способ связи")),
                },
            ))
    return records


def parse_b2c_user_fixture(path: Path) -> list[AdapterRecord]:
    rows = _sheet_rows(path)
    records: list[AdapterRecord] = []
    for index, row in enumerate(rows, start=2):
        full_name = normalize_text(" ".join(str(row.get(key) or "") for key in ("Фамилия", "Имя", "Отчествопри наличии)")))
        email = normalize_email(row.get("Email"))
        phone = normalize_phone(row.get("Номер телефона"))
        external_key = hash_value(email or phone or f"row:{index}")
        records.append(AdapterRecord(
            source="B2C_USER", external_key=external_key or f"row:{index}", raw_payload=row,
            normalized_payload={"external_person_key": external_key, "full_name": full_name, "email_hash": hash_value(email), "phone_hash": hash_value(phone)},
        ))
    return records


def parse_payment_fixture(path: Path) -> list[AdapterRecord]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(payload, list):
        raise ValueError("Payment fixture must contain an array")
    records: list[AdapterRecord] = []
    for item in payload:
        if item is None:
            continue
        if not isinstance(item, dict):
            raise ValueError("Payment fixture item must be an object or null")
        order = normalize_text(item.get("Номер заявки"))
        if not order:
            raise ValueError("Payment fixture item has no order number")
        email = normalize_email(item.get("Email"))
        phone = normalize_phone(item.get("Телефон"))
        full_name = normalize_text(" ".join(str(item.get(key) or "") for key in ("Фамилия", "Имя", "Отчество")))
        course = normalize_text(item.get("Курс"))
        stream = normalize_text(item.get("Номер потока"))
        records.append(AdapterRecord(
            source="PAYMENT", external_key=order, raw_payload=item,
            normalized_payload={"order_number": order, "external_course_name": course, "external_course_key": normalized_key(course),
                                "external_stream_id": stream, "full_name": full_name, "email_hash": hash_value(email), "phone_hash": hash_value(phone)},
        ))
    return records


def parse_team_fixture(source: str, path: Path | None = None) -> list[AdapterRecord]:
    if source not in {"WEBSITE", "LMS"}:
        raise ValueError("Only WEBSITE and LMS have team fixtures")
    records: list[AdapterRecord] = []
    fixture_payload = json.loads((path or TEAM_FIXTURE_PATH).read_text(encoding="utf-8-sig"))
    signals = fixture_payload.get("signals", []) if isinstance(fixture_payload, dict) else fixture_payload
    if not isinstance(signals, list):
        raise ValueError("Website/LMS fixture must contain an array or a signals array")
    for index, item in enumerate(signals, start=1):
        if not isinstance(item, dict):
            continue
        if item.get("source", source).upper() != source:
            continue
        match = item.get("match") or {}
        raw = dict(item)
        normalized = {"organization_name": normalize_text(match.get("organization_name")), "product_name": normalize_text(match.get("product_name")), "organization_code": normalize_text(item.get("organization_code")), "product_code": normalize_text(item.get("product_code")), **(item.get("payload") or {})}
        records.append(AdapterRecord(source=source, external_key=normalize_text(item.get("external_event_id")) or f"team:{source}:{index}", raw_payload=raw, normalized_payload=normalized))
    return records


def utcnow() -> datetime:
    return datetime.now(timezone.utc)
