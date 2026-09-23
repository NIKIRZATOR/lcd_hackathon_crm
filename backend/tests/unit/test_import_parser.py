from io import BytesIO
from pathlib import Path

import pytest
import xlwt
from openpyxl import Workbook

from app.modules.imports.parser.factory import reader_for_path
from app.modules.imports.parser.base import SpreadsheetError


def write_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"
    sheet.append(["Название ВУЗа", "Вендор", "ПО"])
    sheet.append(["МГУ", "Ростелеком", "Продукт 1"])
    workbook.save(path)


def write_xls(path: Path) -> None:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Лист1")
    for col, value in enumerate(["Название ВУЗа", "Вендор", "ПО"]):
        sheet.write(0, col, value)
    for col, value in enumerate(["МГУ", "Ростелеком", "Продукт 1"]):
        sheet.write(1, col, value)
    output = BytesIO()
    workbook.save(output)
    path.write_bytes(output.getvalue())


def test_xlsx_preview_reads_cyrillic_headers_and_rows(tmp_path: Path) -> None:
    path = tmp_path / "catalog.xlsx"
    write_xlsx(path)

    preview = reader_for_path(path, "catalog.xlsx").preview()

    assert preview.file_type == "xlsx"
    assert preview.headers == ["Название ВУЗа", "Вендор", "ПО"]
    assert preview.rows == [["МГУ", "Ростелеком", "Продукт 1"]]


def test_xls_preview_reads_legacy_biff(tmp_path: Path) -> None:
    path = tmp_path / "catalog.xls"
    write_xls(path)

    preview = reader_for_path(path, "catalog.xls").preview()

    assert preview.file_type == "xls"
    assert preview.headers[0] == "Название ВУЗа"
    assert preview.rows[0][0] == "МГУ"


def test_fake_xls_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "fake.xls"
    write_xlsx(path)

    with pytest.raises(SpreadsheetError) as error:
        reader_for_path(path, "fake.xls")

    assert error.value.detail["code"] == "IMPORT_INVALID_FORMAT"


def test_unsupported_extension_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "catalog.csv"
    path.write_text("a,b\n1,2", encoding="utf-8")

    with pytest.raises(SpreadsheetError) as error:
        reader_for_path(path, "catalog.csv")

    assert error.value.detail["code"] == "IMPORT_UNSUPPORTED_FILE"


def test_xlsx_rejects_too_many_columns(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "wide.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["A", "B"])
    sheet.append(["1", "2"])
    workbook.save(path)
    monkeypatch.setattr("app.modules.imports.parser.xlsx.settings.import_max_columns", 1)

    with pytest.raises(SpreadsheetError) as error:
        reader_for_path(path, "wide.xlsx").preview()

    assert error.value.detail["code"] == "IMPORT_TOO_MANY_COLUMNS"
